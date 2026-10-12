"""Transparent, reproducible audit engine for the workplace dust research files.

This module checks consistency and documented derivations. It deliberately does
not certify that observations describe real people or events: that requires
independent primary-source evidence.
"""
from __future__ import annotations

import hashlib
import io
import re
from collections import Counter
from decimal import Decimal, InvalidOperation
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

WORKER_ID_ALIASES = {"workerid", "worker_id", "workeridentifier", "id"}
SYMPTOM_COLUMNS = [
    "cough", "phlegm", "wheezing", "breathlessness", "chest_tightness",
    "eye_irritation", "skin_irritation", "nasal_irritation", "fatigue", "headache",
]
ZONE_ALIASES = {"high": "High", "medium": "Medium", "low": "Low"}
EVIDENCE_TERMS = re.compile(r"simulat|illustrativ|synthetic|fabricat|generated data|artificial data", re.I)


def slug(value: Any) -> str:
    s = str(value or "").strip().lower()
    s = re.sub(r"[^a-z0-9]+", "_", s).strip("_")
    if s in {"workerid", "worker_identifier", "worker_id"}:
        return "worker_id"
    if s == "exposurezone":
        return "exposure_zone"
    return s or "unnamed_column"


def unique_columns(columns) -> list[str]:
    seen: Counter = Counter()
    output = []
    for col in columns:
        name = slug(col)
        seen[name] += 1
        output.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return output


def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = unique_columns(out.columns)
    out = out.dropna(how="all").reset_index(drop=True)
    # Normalize worker IDs and category whitespace but don't silently change numeric values.
    if "worker_id" in out.columns:
        out["worker_id"] = out["worker_id"].map(lambda x: str(x).strip() if pd.notna(x) else np.nan)
    if "exposure_zone" in out.columns:
        out["exposure_zone"] = out["exposure_zone"].map(
            lambda x: ZONE_ALIASES.get(str(x).strip().lower(), str(x).strip()) if pd.notna(x) else np.nan
        )
    return out


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def source_files_from_directory(data_dir) -> list[dict]:
    from pathlib import Path
    files = []
    for path in sorted(Path(data_dir).glob("*")):
        if path.suffix.lower() in {".xlsx", ".xlsm", ".csv"}:
            b = path.read_bytes()
            files.append({"name": path.name, "bytes": b, "origin": "Bundled project file"})
    return files


def _find_worker_header(raw: pd.DataFrame, max_rows: int = 8) -> int | None:
    # Avoid mistaking a Variable Map row containing the literal "Worker_ID" for a data header.
    # A real worker table header should also contain at least one familiar field name.
    known_headers = {
        "age", "job_role", "exposure_zone", "shift_hours", "total_tenure_years",
        "cough", "health_symptom_score", "safety_practice_score", "exposure_index",
        "presenteeism_score", "performance_impairment_score", "smoking_status",
    }
    for idx in range(min(max_rows, len(raw))):
        values = [slug(v) for v in raw.iloc[idx].tolist() if pd.notna(v)]
        normalized = {v.replace("_", "") for v in values}
        known = {v for v in values if v in known_headers}
        if normalized.intersection({"workerid", "workeridentifier"}) and known:
            return idx
    return None


def inspect_source_file(file_obj: dict) -> dict:
    name, data = file_obj["name"], file_obj["bytes"]
    result = {
        "name": name, "origin": file_obj.get("origin", "Uploaded file"), "size_bytes": len(data),
        "sha256": file_hash(data), "sheets": [], "tables": [], "formula_count": 0,
        "formula_samples": [], "evidence_phrases": [], "cell_formats": {}, "load_error": None,
    }
    suffix = name.lower().rsplit(".", 1)[-1]
    try:
        if suffix == "csv":
            df = pd.read_csv(io.BytesIO(data), dtype=object)
            clean = clean_dataframe(df)
            result["sheets"].append({"sheet": "CSV", "rows": len(df), "columns": len(df.columns)})
            if "worker_id" in clean.columns:
                result["tables"].append({"file": name, "sheet": "CSV", "header_row": 0, "df": clean,
                                         "original_columns": list(df.columns), "source_kind": "CSV"})
            for ridx, row in df.iterrows():
                for cidx, value in enumerate(row.tolist()):
                    if isinstance(value, str) and EVIDENCE_TERMS.search(value):
                        result["evidence_phrases"].append({"sheet": "CSV", "cell": f"{get_column_letter(cidx+1)}{ridx+2}", "text": value})
        else:
            wb_formula = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
            for ws in wb_formula.worksheets:
                result["sheets"].append({"sheet": ws.title, "rows": ws.max_row, "columns": ws.max_column})
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            result["formula_count"] += 1
                            if len(result["formula_samples"]) < 12:
                                result["formula_samples"].append({"sheet": ws.title, "cell": cell.coordinate, "formula": cell.value})
                        if isinstance(cell.value, str) and EVIDENCE_TERMS.search(cell.value):
                            result["evidence_phrases"].append({"sheet": ws.title, "cell": cell.coordinate, "text": cell.value})
            wb_formula.close()
            xls = pd.ExcelFile(io.BytesIO(data), engine="openpyxl")
            for sheet in xls.sheet_names:
                raw = pd.read_excel(io.BytesIO(data), sheet_name=sheet, header=None, dtype=object, engine="openpyxl")
                hdr = _find_worker_header(raw)
                if hdr is not None:
                    original = pd.read_excel(io.BytesIO(data), sheet_name=sheet, header=hdr, dtype=object, engine="openpyxl")
                    clean = clean_dataframe(original)
                    if "worker_id" in clean.columns and len(clean) > 0:
                        result["tables"].append({"file": name, "sheet": sheet, "header_row": hdr,
                                                  "df": clean, "original_columns": list(original.columns),
                                                  "source_kind": "Excel"})
            # Capture cell formats lazily only for the few cell traces the UI requests.
        result["evidence_phrases"] = _deduplicate_dicts(result["evidence_phrases"], ("sheet", "cell", "text"))
    except Exception as exc:
        result["load_error"] = f"Could not read this file: {type(exc).__name__}: {exc}"
    return result


def _deduplicate_dicts(items, keys):
    seen, output = set(), []
    for item in items:
        k = tuple(item.get(key) for key in keys)
        if k not in seen:
            seen.add(k)
            output.append(item)
    return output


def analyze_files(file_objects: list[dict]) -> dict:
    # Uploaded versions replace bundled versions with the same filename.
    by_name = {}
    for f in file_objects:
        by_name[f["name"]] = f
    inspected = [inspect_source_file(f) for f in by_name.values()]
    tables = [t for f in inspected for t in f["tables"]]
    evidence = [{**e, "file": f["name"]} for f in inspected for e in f["evidence_phrases"]]
    return {"files": inspected, "tables": tables, "evidence": evidence}


def find_col(df: pd.DataFrame, *candidates: str, prefer_final: bool = False) -> str | None:
    names = list(df.columns)
    for candidate in candidates:
        c = slug(candidate)
        exact = [n for n in names if n == c]
        if exact:
            if prefer_final:
                finals = [n for n in exact if n.endswith("_final")]
                if finals:
                    return finals[0]
            return exact[0]
        starts = [n for n in names if n.startswith(c + "_")]
        if starts:
            if prefer_final:
                finals = [n for n in starts if n.endswith("_final")]
                if finals:
                    return finals[0]
            return starts[0]
    return None


def table_label(table: dict) -> str:
    return f"{table['file']}  ·  {table['sheet']}  ·  {len(table['df'])} rows"


def find_table(tables: list[dict], file_starts: tuple[str, ...] = (), sheet_contains: str = "") -> dict | None:
    matches = []
    for t in tables:
        fn = t["file"].lower()
        sh = t["sheet"].lower()
        if file_starts and not any(fn.startswith(x.lower()) for x in file_starts):
            continue
        if sheet_contains and sheet_contains.lower() not in sh:
            continue
        matches.append(t)
    return matches[0] if matches else None


def key_tables(tables: list[dict]) -> dict[str, dict | None]:
    return {
        "profile": find_table(tables, ("Dataset1",), "worker profile") or find_table(tables, ("D1 Research",), "worker profile"),
        "health": find_table(tables, ("Dataset2",), "health"),
        "safety": find_table(tables, ("Dataset3",), "safety"),
        "performance": find_table(tables, ("Dataset4(",), "performance"),
        "performance_csv": find_table(tables, ("Dataset4_Orange",), ""),
        "master": find_table(tables, ("D5 MASTER",), "master dataset") or find_table(tables, ("Dataset5m",), "master dataset"),
        "analysis": find_table(tables, ("Dataset6",), "analysis dataset"),
    }


def _norm_value(v):
    if pd.isna(v):
        return None
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    text = str(v).strip()
    # Excel, CSV and normalized workbooks may encode the same number with different dtypes.
    try:
        return float(text.replace(",", ""))
    except (ValueError, TypeError):
        return text.lower()


def _equal_values(a, b) -> bool:
    va, vb = _norm_value(a), _norm_value(b)
    if va is None or vb is None:
        return va is vb
    if isinstance(va, float) and isinstance(vb, float):
        return bool(np.isclose(va, vb, atol=1e-9, rtol=0))
    return va == vb


def reconcile_tables(left: dict, right: dict) -> dict:
    a, b = left["df"].copy(), right["df"].copy()
    if "worker_id" not in a.columns or "worker_id" not in b.columns:
        return {"left": table_label(left), "right": table_label(right), "common_ids": 0, "only_left": len(a), "only_right": len(b), "field_results": [], "status": "Not comparable"}
    aid = a["worker_id"].astype(str).str.strip()
    bid = b["worker_id"].astype(str).str.strip()
    a.index, b.index = aid, bid
    common = sorted(set(a.index).intersection(b.index))
    fields = sorted((set(a.columns) & set(b.columns)) - {"worker_id"})
    results = []
    for field in fields:
        matches, compared, missing = 0, 0, 0
        for wid in common:
            av, bv = a.loc[wid, field], b.loc[wid, field]
            if isinstance(av, pd.Series) or isinstance(bv, pd.Series):
                continue
            if pd.isna(av) or pd.isna(bv):
                missing += 1
                continue
            compared += 1
            matches += int(_equal_values(av, bv))
        results.append({"field": field, "matched": matches, "compared": compared, "mismatches": compared-matches, "missing_pairs": missing})
    mismatches = sum(r["mismatches"] for r in results)
    status = "Exact match on shared fields" if mismatches == 0 and len(common) == len(a) == len(b) else ("Differences found" if mismatches else "Partial overlap; shared values match")
    return {"left": table_label(left), "right": table_label(right), "common_ids": len(common),
            "only_left": len(set(a.index)-set(b.index)), "only_right": len(set(b.index)-set(a.index)),
            "field_results": results, "status": status}


def compute_health_score_trace(df: pd.DataFrame, row: pd.Series, target_col: str | None = None) -> dict | None:
    symptom_cols = [c for c in SYMPTOM_COLUMNS if c in df.columns]
    if len(symptom_cols) < 8:
        return None
    vals = []
    for c in symptom_cols:
        v = row.get(c)
        vals.append(v)
    numeric = pd.to_numeric(pd.Series(vals), errors="coerce")
    if numeric.isna().any():
        return {"available": True, "valid": False, "reason": "One or more symptom inputs are missing or non-numeric.", "inputs": dict(zip(symptom_cols, vals))}
    total = float(numeric.sum())
    score_col = target_col if target_col in df.columns else find_col(df, "health_symptom_score")
    observed = row.get(score_col) if score_col else None
    exact = observed is not None and pd.notna(observed) and np.isclose(float(observed), total, atol=1e-9)
    return {"available": True, "valid": True, "inputs": dict(zip(symptom_cols, vals)), "sum": total,
            "score_col": score_col, "observed": observed, "matches": bool(exact)}


def _cell_numeric_format(data: bytes, filename: str, sheet: str, coordinate: str) -> str | None:
    if filename.lower().endswith(".csv"):
        return "CSV has no Excel cell number format."
    try:
        wb = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
        ws = wb[sheet]
        value = ws[coordinate].number_format
        wb.close()
        return value
    except Exception:
        return None


def trace_selected_cell(table: dict, file_lookup: dict[str, dict], row_index: int, column: str) -> dict:
    df = table["df"]
    row = df.iloc[row_index]
    source = file_lookup.get(table["file"], {})
    if column not in df.columns:
        raise KeyError(column)
    col_idx = list(df.columns).index(column) + 1
    header_excel_row = int(table.get("header_row", 0)) + 1
    excel_row = header_excel_row + 1 + row_index
    coordinate = f"{get_column_letter(col_idx)}{excel_row}"
    raw_value = row.get(column)
    numeric = pd.to_numeric(pd.Series([raw_value]), errors="coerce").iloc[0]
    is_numeric = pd.notna(numeric)
    decimals = decimal_places(raw_value) if is_numeric else None
    formula_trace = None
    if column.startswith("health_symptom_score") or column == "health_symptom_score":
        formula_trace = compute_health_score_trace(df, row, target_col=column)
    if column in {"safety_practice_score", "safety_practice_score_final", "safety_practice_score_2"}:
        formula_trace = {"available": False, "reason": "The workbook coding guide describes a score but does not provide the exact weights, transformation, or calculation formula. It cannot be independently recomputed exactly from the workbook documentation."}
    elif column.startswith("ppe_compliance"):
        formula_trace = {"available": False, "reason": "The coding guide calls this an illustrative percentage based mainly on PPE availability/use, but gives no exact formula or weights."}
    elif column.startswith("performance_impairment_score"):
        formula_trace = {"available": False, "reason": "The coding guide calls this an illustrative composite score but does not provide component weights or the exact formula."}
    elif column == "exposure_index":
        formula_trace = {"available": False, "reason": "The workbook labels this as derived but does not specify an exact calculation formula or component weights."}
    elif formula_trace is None:
        formula_trace = {"available": False, "reason": "No documented cell formula or explicit derivation rule was found for this field. The app reports the stored value and its location rather than guessing a formula."}
    fmt = _cell_numeric_format(source.get("bytes", b""), table["file"], table["sheet"], coordinate) if source else None
    return {"file": table["file"], "sheet": table["sheet"], "worker_id": row.get("worker_id"),
            "field": column, "value": raw_value, "is_numeric": bool(is_numeric), "decimal_places": decimals,
            "coordinate": coordinate, "excel_number_format": fmt, "formula_trace": formula_trace,
            "row": row.to_dict()}


def decimal_places(value: Any) -> int:
    try:
        d = Decimal(str(value)).normalize()
        return max(0, -d.as_tuple().exponent)
    except (InvalidOperation, ValueError, TypeError):
        return 0


def numeric_decimal_summary(series: pd.Series) -> dict:
    vals = pd.to_numeric(series, errors="coerce").dropna()
    dec_vals = [v for v in vals if not float(v).is_integer()]
    places = [decimal_places(v) for v in dec_vals]
    return {"numeric_count": int(vals.size), "decimal_count": len(dec_vals),
            "integer_count": int(vals.size-len(dec_vals)),
            "max_decimal_places": max(places, default=0),
            "decimal_place_distribution": dict(sorted(Counter(places).items())),
            "min": float(vals.min()) if not vals.empty else None,
            "max": float(vals.max()) if not vals.empty else None}


def run_integrity_checks(tables: list[dict], inspected_files: list[dict]) -> list[dict]:
    checks = []
    def add(area, check, status, finding, detail=""):
        checks.append({"area": area, "check": check, "status": status, "finding": finding, "detail": detail})
    valid_tables = [t for t in tables if not t.get("df", pd.DataFrame()).empty]
    add("File integrity", "SHA-256 fingerprints recorded", "PASS" if inspected_files else "WARN",
        f"{len(inspected_files)} source files inspected.", "Fingerprints identify the exact bytes inspected; they do not prove who created the files or whether the observations are true.")
    formula_total = sum(f.get("formula_count", 0) for f in inspected_files)
    add("Formula lineage", "Spreadsheet formulas present", "PASS" if formula_total else "WARN",
        f"Found {formula_total} formula cells across the Excel workbooks.", "If zero, calculations are stored as values rather than live Excel formulas. A written formula specification is still needed to validate derived scores.")
    evidence_count = sum(len(f.get("evidence_phrases", [])) for f in inspected_files)
    evidence_status = "WARN" if evidence_count else "INFO"
    add("Provenance", "Simulation / illustrative wording", evidence_status,
        f"Found {evidence_count} wording matches across files.", "These phrases are shown verbatim in the Evidence tab. Wording is evidence about how the workbook describes the data, not a cryptographic proof of origin.")
    for t in valid_tables:
        df = t["df"]
        label = f"{t['file']} / {t['sheet']}"
        if "worker_id" not in df.columns:
            continue
        ids = df["worker_id"].dropna().astype(str).str.strip()
        duplicate_ids = int(ids.duplicated().sum())
        missing_ids = int(df["worker_id"].isna().sum())
        add("Structure", f"Unique worker IDs — {label}", "PASS" if duplicate_ids == 0 and missing_ids == 0 else "FAIL",
            f"{len(ids)} nonblank IDs; {duplicate_ids} duplicate IDs; {missing_ids} missing IDs.")
        zone_col = find_col(df, "exposure_zone")
        if zone_col:
            counts = df[zone_col].value_counts(dropna=False).to_dict()
            expected = {"High": 35, "Medium": 25, "Low": 20}
            zone_valid = counts == expected
            # Use only exact check for 80-person profile/master data; related sheets should share the assignment.
            add("Structure", f"Exposure allocation — {label}", "PASS" if zone_valid else "WARN",
                f"Observed zone counts: {counts}.", "Project target allocation is High 35, Medium 25, Low 20; a mismatch can indicate a different subset or a data issue.")
    roles = key_tables(tables)
    master = roles.get("master") or roles.get("analysis")
    if master:
        df = master["df"]
        n = len(df)
        add("Structure", "Master/analysis table has 80 records", "PASS" if n == 80 else "WARN",
            f"Selected integrated table has {n} rows.")
        for label, candidates in [("Exposure_Index", ("exposure_index",)),
                                  ("Health symptom score", ("health_symptom_score_final", "health_symptom_score")),
                                  ("Safety practice score", ("safety_practice_score_final", "safety_practice_score")),
                                  ("Performance impairment score", ("performance_impairment_score_final", "performance_impairment_score"))]:
            col = find_col(df, *candidates, prefer_final=True)
            if col:
                numeric = pd.to_numeric(df[col], errors="coerce")
                out_of_range = 0
                if "health" in label.lower(): out_of_range = int(((numeric < 0) | (numeric > 10)).sum())
                elif label == "Exposure_Index": out_of_range = int(((numeric < 0) | (numeric > 100)).sum())
                elif label == "Safety practice score": out_of_range = int(((numeric < 0) | (numeric > 100)).sum())
                add("Range", f"{label} plausible range", "PASS" if out_of_range == 0 else "FAIL",
                    f"Column '{col}': {out_of_range} values outside tested range.", "Range checks detect impossible or unexpected values under the stated scale; they do not prove observed data authenticity.")
        symptoms = [c for c in SYMPTOM_COLUMNS if c in df.columns]
        if len(symptoms) == 10:
            symptom_values = df[symptoms].apply(pd.to_numeric, errors="coerce")
            invalid_binary = int(((~symptom_values.isin([0, 1])) & symptom_values.notna()).sum().sum())
            add("Derived value", "Health score inputs are binary", "PASS" if invalid_binary == 0 else "FAIL",
                f"Found {invalid_binary} symptom input values outside 0/1.")
            expected_sum = symptom_values.sum(axis=1, min_count=10)
            score_candidates = [c for c in df.columns if c.startswith("health_symptom_score")]
            for c in score_candidates:
                observed = pd.to_numeric(df[c], errors="coerce")
                compared = expected_sum.notna() & observed.notna()
                mismatches = int((~np.isclose(expected_sum[compared], observed[compared], atol=1e-9)).sum()) if compared.any() else 0
                add("Derived value", f"Health score reproduces from 10 symptoms ({c})", "PASS" if compared.any() and mismatches == 0 else ("FAIL" if mismatches else "WARN"),
                    f"{int(compared.sum())} comparable rows; {mismatches} mismatch(es).",
                    "Documented formula: add the 10 binary symptom indicators. A match proves calculation consistency, not that the symptoms were reported by real workers.")
    # Reconcile principal source tables with integrated master and each other.
    if master:
        for role in ["profile", "health", "safety", "performance", "performance_csv"]:
            t = roles.get(role)
            if t and t is not master:
                recon = reconcile_tables(t, master)
                add("Cross-file reconciliation", f"{role.title()} vs integrated table", "PASS" if recon["status"].startswith("Exact") or recon["status"].startswith("Partial") else "WARN",
                    f"{recon['status']}; common IDs: {recon['common_ids']}; fields compared: {len(recon['field_results'])}.",
                    f"IDs only in source: {recon['only_left']}; only in integrated table: {recon['only_right']}. Shared-column mismatches: {sum(x['mismatches'] for x in recon['field_results'])}.")
    return checks


def recompute_zone_summary(tables: list[dict]) -> pd.DataFrame:
    roles = key_tables(tables)
    base = roles.get("master") or roles.get("analysis")
    if not base:
        return pd.DataFrame()
    df = base["df"]
    zone = find_col(df, "exposure_zone")
    if not zone:
        return pd.DataFrame()
    metrics = {
        "Mean Exposure Index": find_col(df, "exposure_index"),
        "Mean Health Score": find_col(df, "health_symptom_score_final", "health_symptom_score", prefer_final=True),
        "Mean Safety Score": find_col(df, "safety_practice_score_final", "safety_practice_score", prefer_final=True),
        "Mean PPE Compliance": find_col(df, "ppe_compliance"),
        "Mean Performance Impairment": find_col(df, "performance_impairment_score_final", "performance_impairment_score", prefer_final=True),
    }
    out = df.groupby(zone, dropna=False).size().rename("N").to_frame()
    for label, col in metrics.items():
        if col:
            out[label] = df.groupby(zone)[col].apply(lambda s: pd.to_numeric(s, errors="coerce").mean())
    out.index.name = "Exposure Zone"
    out = out.reset_index()
    return out


def correlation_summary(tables: list[dict]) -> pd.DataFrame:
    roles = key_tables(tables)
    base = roles.get("master") or roles.get("analysis")
    if not base:
        return pd.DataFrame()
    df = base["df"]
    cols = {
        "Exposure Index": find_col(df, "exposure_index"),
        "Health Symptoms": find_col(df, "health_symptom_score_final", "health_symptom_score", prefer_final=True),
        "Safety Practice": find_col(df, "safety_practice_score_final", "safety_practice_score", prefer_final=True),
        "PPE Compliance": find_col(df, "ppe_compliance"),
        "Performance Impairment": find_col(df, "performance_impairment_score_final", "performance_impairment_score", prefer_final=True),
    }
    cols = {label: col for label, col in cols.items() if col}
    if len(cols) < 2:
        return pd.DataFrame()
    num = pd.DataFrame({label: pd.to_numeric(df[col], errors="coerce") for label, col in cols.items()})
    return num.corr(method="pearson").round(3)


def report_markdown(analysis: dict, tables: list[dict]) -> str:
    files = analysis.get("files", [])
    checks = run_integrity_checks(tables, files)
    output = [
        "# Research Data Audit Report",
        "",
        "## Scope and interpretation",
        "This report checks file integrity, structural consistency, cross-file agreement, and whether documented derived values reproduce. It does **not** certify real-world authenticity from spreadsheets alone.",
        "",
        f"- Source files inspected: {len(files)}",
        f"- Worker-level tables extracted: {len(tables)}",
        f"- Simulation/illustrative wording hits: {len(analysis.get('evidence', []))}",
        f"- Spreadsheet formula cells: {sum(f.get('formula_count', 0) for f in files)}",
        "",
        "## Key provenance limitation",
        "At least one workbook explicitly describes a result as observed in 'simulated data'; several codebooks use 'illustrative' language for variables. Any claim that the observations are actual worker responses requires the original field records, data collection process, and exact score formulas.",
        "",
        "## Integrity checks",
        "| Area | Check | Status | Finding |",
        "|---|---|---|---|",
    ]
    for c in checks:
        output.append(f"| {c['area']} | {c['check']} | {c['status']} | {c['finding'].replace('|', '/') } |")
    output.extend(["", "## Evidence phrases"])
    if analysis.get("evidence"):
        output.extend(["| File | Sheet | Cell | Wording |", "|---|---|---|---|"])
        for e in analysis["evidence"]:
            text = str(e["text"]).replace("|", "/").replace("\n", " ")
            output.append(f"| {e['file']} | {e['sheet']} | {e['cell']} | {text} |")
    else:
        output.append("No explicit simulation/illustrative wording found by the current phrase scan. Absence of a phrase is not proof of authenticity.")
    output.extend(["", "## Recomputed zone summary"])
    zone = recompute_zone_summary(tables)
    if not zone.empty:
        output.append(zone.to_markdown(index=False))
    else:
        output.append("No integrated worker-level table was identified.")
    output.extend(["", "## Decimal interpretation", "A decimal is a representation of a value, not proof of provenance. A value such as 39.3 can be (a) a raw entry, (b) a transformed/composite score, or (c) a summary such as a mean or correlation. Only values with a documented formula and underlying inputs can be independently recomputed; otherwise this report records the stored value and flags missing formula/source documentation.", ""])
    return "\n".join(output)
